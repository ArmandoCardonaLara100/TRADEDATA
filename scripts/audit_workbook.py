"""Read-only, reproducible audit of every populated cell and native chart."""
import json, pathlib, sys, hashlib, collections
import openpyxl
from openpyxl.formula.translate import Translator

source = pathlib.Path(sys.argv[1])
out = pathlib.Path(__file__).resolve().parents[1] / 'docs' / 'spreadsheet'
out.mkdir(parents=True, exist_ok=True)
wb = openpyxl.load_workbook(source, data_only=False)
cached = openpyxl.load_workbook(source, data_only=True)
audit = {'source': source.name, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'definedNames': {k: str(v) for k,v in wb.defined_names.items()}, 'sheets': []}
for ws in wb:
    cells, families = [], collections.defaultdict(list)
    for row in ws:
        for c in row:
            if c.value is None: continue
            value = c.value
            formula = value if c.data_type == 'f' else None
            if c.data_type == 'f' and not isinstance(formula, str): formula = getattr(value, 'text', str(value))
            entry = {'cell': c.coordinate, 'value': formula or value, 'cached': cached[ws.title][c.coordinate].value, 'type': c.data_type, 'format': c.number_format}
            if formula:
                try: normalized = Translator(formula, origin=c.coordinate).translate_formula('XFD100000')
                except Exception: normalized = formula
                families[normalized].append(c.coordinate)
            cells.append(entry)
    charts = []
    for chart in ws._charts:
        charts.append({'type':type(chart).__name__, 'xml':openpyxl.xml.functions.tostring(chart.to_tree()).decode()})
    audit['sheets'].append({'name': ws.title, 'state':ws.sheet_state, 'dimensions':ws.calculate_dimension(), 'cells':cells, 'formulaFamilies':[{'formula':cells[next(i for i,x in enumerate(cells) if x['cell']==refs[0])]['value'], 'cells':refs} for f,refs in families.items()], 'charts':charts, 'tables':[str(t) for t in ws.tables.values()], 'validations':[str(v) for v in ws.data_validations.dataValidation], 'conditionalFormatting':[{'range':str(k),'rules':[str(r) for r in v]} for k,v in ws.conditional_formatting._cf_rules.items()], 'mergedRanges':[str(x) for x in ws.merged_cells.ranges]})
(out/'workbook-audit.json').write_text(json.dumps(audit, default=str, indent=2))
lines=[]
for sheet in audit['sheets']:
    lines += ['\n## '+sheet['name']+' '+sheet['dimensions'], f"{len(sheet['cells'])} nonempty cells, {sum(len(f['cells']) for f in sheet['formulaFamilies'])} formulas, {len(sheet['formulaFamilies'])} families, {len(sheet['charts'])} charts"]
    lines += [f"{c['cell']}: {c['value']}" for c in sheet['cells'] if c['type']!='f']
    lines += ['FORMULA FAMILIES:']
    lines += [f"{f['cells'][0]}..{f['cells'][-1]} ({len(f['cells'])}): {f['formula']} => cached {cached[sheet['name']][f['cells'][0]].value}" for f in sheet['formulaFamilies']]
(out/'inventory.txt').write_text('\n'.join(lines))
print('\n'.join(f"{s['name']}: {len(s['cells'])} cells, {sum(len(f['cells']) for f in s['formulaFamilies'])} formulas, {len(s['formulaFamilies'])} families, {[c['type'] for c in s['charts']]}" for s in audit['sheets']))
